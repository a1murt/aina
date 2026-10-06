"""Reference implementation for the golden values of the case import.

Reads data/case/source/case2_data.docx, writes CSV/XLSX fixtures and
data/case/expected/import_expected.json.

This script is the SOURCE OF TRUTH for golden tests (SPEC.md §12, §14).
The production import + KPI + DQ + rules code must reproduce import_expected.json
(floats within 1e-4). If a default in config/*.yaml changes, update the constants
below, rerun this script and commit the regenerated JSON together with the change.

Run:  python -I data/case/expected/compute_expected.py  (from repo root)
"""
from __future__ import annotations

import csv
import json
import math
import re
import sys
from collections import defaultdict
from datetime import date
from pathlib import Path

import docx  # python-docx
import openpyxl

ROOT = Path(__file__).resolve().parents[3]
SRC = ROOT / "data/case/source/case2_data.docx"
CSV_DIR = ROOT / "data/case/csv"
OUT_JSON = ROOT / "data/case/expected/import_expected.json"
OUT_XLSX = ROOT / "data/case/csv/case2_data.xlsx"

# ---- defaults mirrored from config/*.yaml (keep in sync) -------------------
SHIFT_MIN = 480                      # plant.yaml shifts[*] length (8 h)
ICT_S = 233                          # plant.yaml lines[WELD-1|PAINT-1|ASSY-1].ict_seconds
OEE_TARGET = 0.85                    # rules.yaml thresholds.oee_target
OEE_NEAR_PP = 3.0                    # rules.yaml thresholds.oee_near_margin_pp
DEFECT_LIMIT = 0.02                  # rules.yaml thresholds.defect_rate_limit
DEFECT_CRIT = 0.04                   # rules.yaml thresholds.defect_rate_critical
CRIT_DOWN_LIMIT = 60                 # rules.yaml thresholds.critical_downtime_limit_min_per_day
CRIT_DOWN_WARN = 0.75                # rules.yaml thresholds.critical_downtime_warn_ratio
DQ_LOAD_PP = 0.5                     # rules.yaml data_quality.load_mismatch_pp
DQ_RECON_MIN = 10                    # rules.yaml data_quality.downtime_recon_min
DQ_FLOW_WARN = 5                     # rules.yaml data_quality.flow_balance_warn_units
WORKING_DAYS_OCT_2026 = 21           # plant.yaml calendar (22 weekdays − 26.10 transferred holiday)
SHIFTS_PER_DAY = 2

AREA = {"Сварка": "WELD", "Окраска": "PAINT", "Сборка": "ASSY"}
LINE = {"Сварка-1": "WELD-1", "Окраска-1": "PAINT-1", "Сборка-1": "ASSY-1"}
LINE_AREA = {"WELD-1": "WELD", "PAINT-1": "PAINT", "ASSY-1": "ASSY"}
FLOW = ["WELD", "PAINT", "ASSY"]
EQUIP = {"ABB-01": "ABB-01", "ABB-04": "ABB-04", "Камера-02": "BOOTH-02", "Конвейер-03": "CONV-03"}
EQUIP_CLASS = {"ABB-01": ("B", 0.5), "ABB-04": ("B", 0.5), "BOOTH-02": ("A", 0.0), "CONV-03": ("A", 0.0)}
REASON = {"Ошибка датчика": ("EL-SENSOR", False), "Замена фильтра": ("MT-FILTER", False),
          "Обрыв цепи": ("ME-CHAIN", False), "Плановое ТО": ("PM-SCHEDULED", True)}
MODEL = {"Chevrolet Onix": "ONIX", "Chevrolet Cobalt": "COBALT", "JAC J7": "J7"}


def num(s: str) -> float:
    s = s.strip().replace(" ", "").replace(" ", "").replace(",", ".")
    return float(s)


def ddmmyyyy(s: str) -> str:
    d, m, y = (int(x) for x in s.strip().split("."))
    return date(y, m, d).isoformat()


def read_tables(path: Path) -> dict[str, list[list[str]]]:
    d = docx.Document(str(path))
    found: dict[str, list[list[str]]] = {}
    for t in d.tables:
        rows = [[c.text.strip() for c in r.cells] for r in t.rows]
        header = " | ".join(rows[0]).lower()
        if "линия" in header and "факт" in header:
            found["lines"] = rows
        elif "оборудование" in header and "причина" in header:
            found["downtime"] = rows
        elif "модель" in header and "план" in header:
            found["plan"] = rows
        elif "выпущено" in header and "брак" in header:
            found["quality"] = rows
    missing = {"lines", "downtime", "plan", "quality"} - found.keys()
    if missing:
        sys.exit(f"tables not found: {missing}")
    found["_text"] = [[p.text] for p in d.paragraphs if p.text.strip()]
    return found


def parse_constraints(paragraphs: list[str]) -> dict:
    txt = "\n".join(paragraphs)
    out = {}
    if m := re.search(r"(\d+)\s*смен[ыа]?\s*по\s*(\d+)\s*час", txt):
        out["shifts_per_day"], out["shift_hours"] = int(m[1]), int(m[2])
    if m := re.search(r"OEE\s*-?\s*не менее\s*(\d+)\s*%", txt):
        out["oee_target"] = int(m[1]) / 100
    if m := re.search(r"брака\s*-?\s*не более\s*(\d+)\s*%", txt):
        out["defect_rate_limit"] = int(m[1]) / 100
    if m := re.search(r"простой критического оборудования\s*-?\s*(\d+)\s*минут", txt):
        out["critical_downtime_limit_min_per_day"] = int(m[1])
    if m := re.search(r"не менее\s*([\d\s ]+)\s*автомобил", txt):
        out["plant_target_per_month"] = int(re.sub(r"\D", "", m[1]))
    return out


def write_csv(name: str, rows: list[list[str]]) -> None:
    CSV_DIR.mkdir(parents=True, exist_ok=True)
    with open(CSV_DIR / name, "w", newline="", encoding="utf-8") as f:
        csv.writer(f, delimiter=";").writerows(rows)


def r4(x: float) -> float:
    return round(x, 4)


def main() -> None:
    t = read_tables(SRC)
    write_csv("01_lines.csv", t["lines"])
    write_csv("02_downtime.csv", t["downtime"])
    write_csv("03_plan.csv", t["plan"])
    write_csv("04_quality.csv", t["quality"])
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet, key in [("Линии", "lines"), ("Простои", "downtime"), ("План", "plan"), ("Качество", "quality")]:
        ws = wb.create_sheet(sheet)
        for row in t[key]:
            ws.append(row)
    wb.save(OUT_XLSX)

    constraints = parse_constraints([p[0] for p in t["_text"]])

    # ---- shift reports ------------------------------------------------------
    quality = {}
    for d, area, out_, rej, pct in t["quality"][1:]:
        quality[(ddmmyyyy(d), AREA[area])] = (int(num(out_)), int(num(rej)), num(pct))

    reports = []
    for d, line, plan, fact, hours, load in t["lines"][1:]:
        day, lc = ddmmyyyy(d), LINE[line]
        area = LINE_AREA[lc]
        produced, worked_min = int(num(fact)), num(hours) * 60
        q_out, defects, pct_src = quality[(day, area)]
        assert q_out == produced, (day, area)
        good = produced - defects
        pbt, apt = SHIFT_MIN, worked_min
        a = apt / pbt
        e = ICT_S * produced / (apt * 60)
        qr = good / produced
        oee = good * ICT_S / (pbt * 60)
        assert math.isclose(a * e * qr, oee, rel_tol=1e-12)
        reports.append({
            "date": day, "shift": "A", "line": lc, "area": area,
            "plan_qty": int(num(plan)), "produced_qty": produced, "defect_qty": defects, "good_qty": good,
            "worked_min": r4(worked_min), "reported_load_pct": num(load), "reported_defect_pct": pct_src,
            "pot_min": SHIFT_MIN, "pdot_min": 0, "pbt_min": pbt, "apt_min": r4(apt), "lost_min": r4(pbt - apt),
            "availability": r4(a), "effectiveness": r4(e), "quality_ratio": r4(qr), "oee": r4(oee),
            "defect_rate": r4(defects / produced), "fpy": r4(qr),
        })

    # ---- downtime -----------------------------------------------------------
    downtime = []
    for d, area, eq, reason, dur in t["downtime"][1:]:
        code, planned = REASON[reason]
        ec = EQUIP[eq]
        cls, degraded = EQUIP_CLASS[ec]
        dur_min = num(dur)
        eff_loss_min = 0.0 if planned else dur_min * (1 - degraded)
        downtime.append({
            "date": ddmmyyyy(d), "area": AREA[area], "equipment": ec, "reason_code": code,
            "reason_text_src": reason, "planned": planned, "duration_min": dur_min, "shift": None,
            "criticality": cls, "degraded_capacity": degraded,
            "effective_capacity_loss_min": r4(eff_loss_min),
            "capacity_loss_units": r4(eff_loss_min * 60 / ICT_S),
        })

    # ---- plan -----------------------------------------------------------------
    plan_rows = [{"model": MODEL[m], "model_src": m, "qty": int(num(q))} for m, q in t["plan"][1:]]
    line_plan = sum(p["qty"] for p in plan_rows)
    target = constraints.get("plant_target_per_month", 5500)
    shifts_month = WORKING_DAYS_OCT_2026 * SHIFTS_PER_DAY
    by_day = defaultdict(dict)
    for r in reports:
        by_day[r["date"]][r["area"]] = r
    days = sorted(by_day)
    bottleneck_by_day = {dd: min(FLOW, key=lambda ar: by_day[dd][ar]["produced_qty"]) for dd in days}
    mean_rate = {ar: sum(by_day[dd][ar]["produced_qty"] for dd in days) / len(days) for ar in FLOW}
    overall_bn = min(FLOW, key=lambda ar: mean_rate[ar])
    mean_min_rate = sum(min(by_day[dd][ar]["produced_qty"] for ar in FLOW) for dd in days) / len(days)

    # ---- data quality ------------------------------------------------------------
    dq = []
    for r in reports:
        diff = r["reported_load_pct"] - round(r["availability"] * 100, 1)
        if abs(diff) > DQ_LOAD_PP:
            dq.append({"rule_id": "DQ-01", "severity": "info", "entity": r["line"], "date": r["date"],
                       "details": {"reported_load_pct": r["reported_load_pct"],
                                   "computed_availability_pct": round(r["availability"] * 100, 1),
                                   "diff_pp": round(diff, 1)}})
    logged = defaultdict(float)
    for x in downtime:
        logged[(x["date"], x["area"])] += x["duration_min"]
    for r in reports:
        lg, lost = logged.get((r["date"], r["area"]), 0.0), r["lost_min"]
        if lg == 0 and lost == 0:
            continue
        diff = lg - lost
        if abs(diff) >= DQ_RECON_MIN:
            dq.append({"rule_id": "DQ-02", "severity": "warning", "entity": r["area"], "date": r["date"],
                       "details": {"logged_downtime_min": lg, "lost_time_min": lost, "diff_min": r4(diff),
                                   "direction": "log_exceeds_loss" if diff > 0 else "unlogged_loss"}})
    no_shift = sum(1 for x in downtime if x["shift"] is None)
    if no_shift:
        dq.append({"rule_id": "DQ-03", "severity": "info", "entity": "downtime_log", "date": None,
                   "details": {"records_without_shift": no_shift}})
    for up, dn in zip(FLOW, FLOW[1:]):
        s_up = sum(by_day[dd][up]["produced_qty"] for dd in days)
        s_dn = sum(by_day[dd][dn]["produced_qty"] for dd in days)
        delta = s_up - s_dn
        if delta < 0:
            dq.append({"rule_id": "DQ-04", "severity": "warning" if abs(delta) >= DQ_FLOW_WARN else "info",
                       "entity": f"{up}->{dn}", "date": None,
                       "details": {"upstream_produced": s_up, "downstream_produced": s_dn,
                                   "buffer_change_units": delta}})
    if line_plan != target:
        dq.append({"rule_id": "DQ-05", "severity": "warning", "entity": "production_plan", "date": None,
                   "details": {"line_model_plan": line_plan, "plant_target": target, "gap": line_plan - target}})
    for r in reports:
        if r["effectiveness"] > 1.0:
            dq.append({"rule_id": "DQ-06", "severity": "warning", "entity": r["line"], "date": r["date"],
                       "details": {"effectiveness": r["effectiveness"]}})

    # ---- alerts ---------------------------------------------------------------------
    alerts = []
    for r in reports:
        dr = r["defect_rate"]
        if dr > DEFECT_LIMIT:
            alerts.append({"rule_id": "AL-Q1", "severity": "critical" if dr > DEFECT_CRIT else "warning",
                           "entity": r["area"], "date": r["date"], "value": dr})
        if r["oee"] < OEE_TARGET:
            alerts.append({"rule_id": "AL-O1", "severity": "warning", "entity": r["line"], "date": r["date"], "value": r["oee"]})
        elif r["oee"] < OEE_TARGET + OEE_NEAR_PP / 100:
            alerts.append({"rule_id": "AL-O2", "severity": "info", "entity": r["line"], "date": r["date"], "value": r["oee"]})
    unplanned_a = defaultdict(float)
    for x in downtime:
        if x["criticality"] == "A" and not x["planned"]:
            unplanned_a[(x["date"], x["equipment"])] += x["duration_min"]
    for (dd, eq), m in sorted(unplanned_a.items()):
        if m >= CRIT_DOWN_LIMIT:
            alerts.append({"rule_id": "AL-D1", "severity": "critical", "entity": eq, "date": dd, "value": m})
        elif m >= CRIT_DOWN_LIMIT * CRIT_DOWN_WARN:
            alerts.append({"rule_id": "AL-D1", "severity": "warning", "entity": eq, "date": dd, "value": m})
    for prev, cur in zip(days, days[1:]):
        up_all = all(by_day[cur][ar]["defect_rate"] > by_day[prev][ar]["defect_rate"] for ar in FLOW)
        any_over = any(by_day[cur][ar]["defect_rate"] > DEFECT_LIMIT for ar in FLOW)
        if up_all and any_over:
            alerts.append({"rule_id": "AL-Q3", "severity": "warning", "entity": "PLANT", "date": cur,
                           "value": {ar: by_day[cur][ar]["defect_rate"] for ar in FLOW}})

    bdr_s = min(r["apt_min"] * 60 / r["produced_qty"] for r in reports)
    out = {
        "meta": {
            "source": "data/case/source/case2_data.docx",
            "ict_seconds": ICT_S, "ict_seconds_bdr": r4(bdr_s), "shift_minutes": SHIFT_MIN,
            "assumptions": ["line table = one shift (A)", "downtime log is per day, shift unknown",
                            "defects are rework (produced counts include defective units)",
                            "pdot=0 for imported shifts because planned downtime cannot be attributed to a shift"],
            "float_tolerance": 1e-4,
        },
        "constraints_parsed": constraints,
        "shift_reports": reports,
        "downtime": downtime,
        "plan": {
            "rows": plan_rows, "line_model_total": line_plan, "plant_target": target, "gap": line_plan - target,
            "working_days_oct_2026": WORKING_DAYS_OCT_2026, "shifts_oct_2026": shifts_month,
            "required_rate_per_shift_line_plan": r4(line_plan / shifts_month),
            "required_rate_per_shift_target": r4(target / shifts_month),
            "mean_sustainable_rate_per_shift": r4(mean_min_rate),
            "naive_month_projection": round(mean_min_rate * shifts_month),
        },
        "flow": {
            "produced_total": {ar: sum(by_day[dd][ar]["produced_qty"] for dd in days) for ar in FLOW},
            "good_total": {ar: sum(by_day[dd][ar]["good_qty"] for dd in days) for ar in FLOW},
            "mean_produced_per_shift": {ar: r4(v) for ar, v in mean_rate.items()},
        },
        "bottleneck_aggregate": {"by_day": bottleneck_by_day, "overall": overall_bn,
                                 "shifting": len(set(bottleneck_by_day.values())) > 1},
        "data_quality_issues": dq,
        "alerts": alerts,
    }
    OUT_JSON.write_text(json.dumps(out, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"ok: {len(reports)} shift reports, {len(downtime)} downtime, {len(dq)} DQ issues, {len(alerts)} alerts")


if __name__ == "__main__":
    main()
