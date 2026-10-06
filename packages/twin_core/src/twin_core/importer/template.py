"""Import template ``template.xlsx`` (FR-IMP-07), generated from the configuration.

Sheets «Линии», «Простои», «План», «Качество» carry the headers the recognizer expects
(SPEC §7.4); «Вводные» holds the constraint sentences with the configured values (edit them to
state the plant's own constraints, FR-IMP-03); «Справочник» lists the names the import accepts.
"""

from __future__ import annotations

import io
from typing import Final

import openpyxl
from openpyxl.styles import Font

from twin_core.config import TwinConfig
from twin_core.importer.constraints import configured_constraints

SHEETS: Final[tuple[tuple[str, tuple[str, ...]], ...]] = (
    ("Линии", ("Дата", "Линия", "План", "Факт", "Время работы, ч", "Загрузка, %")),
    ("Простои", ("Дата", "Участок", "Оборудование", "Причина", "Длительность, мин")),
    ("План", ("Модель", "План на месяц")),
    ("Качество", ("Дата", "Участок", "Выпущено", "Брак", "% брака")),
)
NOTES_SHEET: Final = "Вводные"
REFERENCE_SHEET: Final = "Справочник"


def _number(value: float) -> str:
    text = f"{value:g}" if not float(value).is_integer() else str(int(value))
    return text.replace(".", ",")


def _thousands(value: int) -> str:
    return f"{value:,}".replace(",", " ")


def constraint_sentences(cfg: TwinConfig, month: str | None = None) -> list[str]:
    c = configured_constraints(cfg, month)
    return [
        f"Производство работает в {c['shifts_per_day']} смены "
        f"по {_number(c['shift_hours'])} часов.",
        f"Целевой показатель OEE - не менее {_number(c['oee_target'] * 100)}%.",
        f"Допустимый уровень брака - не более {_number(c['defect_rate_limit'] * 100)}%.",
        "Максимально допустимый простой критического оборудования - "
        f"{_number(c['critical_downtime_limit_min_per_day'])} минут в сутки.",
        "План выпуска - не менее "
        f"{_thousands(int(c['plant_target_per_month']))} автомобилей в месяц.",
    ]


def build_template(cfg: TwinConfig) -> bytes:
    """The xlsx template as bytes."""
    workbook = openpyxl.Workbook()
    default = workbook.active
    if default is not None:
        workbook.remove(default)
    bold = Font(bold=True)
    for title, headers in SHEETS:
        sheet = workbook.create_sheet(title)
        sheet.append(list(headers))
        for column, header in enumerate(headers, start=1):
            sheet.cell(row=1, column=column).font = bold
            letter = sheet.cell(row=1, column=column).column_letter
            sheet.column_dimensions[letter].width = max(14, len(header) + 4)

    notes = workbook.create_sheet(NOTES_SHEET)
    for sentence in constraint_sentences(cfg):
        notes.append([sentence])
    notes.column_dimensions["A"].width = 90

    reference = workbook.create_sheet(REFERENCE_SHEET)
    reference.append(["Вид", "Код", "Название"])
    for column in range(1, 4):
        reference.cell(row=1, column=column).font = bold
    for line_code, line in cfg.lines.items():
        reference.append(["линия", line_code, line.name_ru])
    for area in cfg.areas.values():
        if area.lines:
            reference.append(["участок", area.code, area.name_ru])
    for eq_code, eq in cfg.equipment.items():
        reference.append(["оборудование", eq_code, eq.aliases[0] if eq.aliases else eq.name_ru])
    for reason_code, reason in cfg.reasons.items():
        reference.append(["причина", reason_code, reason.name_ru])
    for product_code, product in cfg.products.items():
        reference.append(["модель", product_code, product.name])
    for letter, width in (("A", 14), ("B", 18), ("C", 40)):
        reference.column_dimensions[letter].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
