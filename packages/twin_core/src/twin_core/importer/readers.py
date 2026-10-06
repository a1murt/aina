"""Format readers: ``.docx``, ``.xlsx``, ``.csv`` / ``.txt`` (or a ``.zip`` of them) -> RawUpload.

Accepted uploads (SPEC §7.4):

* one ``.docx`` — every table plus the body paragraphs as free text;
* one ``.xlsx`` — every sheet as a table (sheet names are informative, recognition is by headers);
  a sheet that is not a recognized table (e.g. «Вводные») is used as free text;
* one or more ``.csv`` (``;``-separated, UTF-8, one table per file) and optional ``.txt`` files
  with free text, uploaded together or packed into one ``.zip``.

Cell values become strings exactly as a user sees them; numbers and dates are parsed later
(:mod:`twin_core.importer.values`) so that every format goes through the same rules.
"""

from __future__ import annotations

import csv
import io
import zipfile
from collections.abc import Iterable, Sequence
from datetime import date, datetime, time
from pathlib import PurePosixPath
from typing import Final

import docx
import openpyxl

from twin_core.importer.model import (
    ImportFormatError,
    ImportKind,
    RawTable,
    RawUpload,
    UploadedFile,
)

CSV_DELIMITER: Final = ";"
MAX_ZIP_MEMBERS: Final = 50
MAX_UNPACKED_BYTES: Final = 20 * 1024 * 1024
"""Upper bound for the unpacked size of a ``.zip`` upload (zip-bomb guard)."""

_TEXT_ENCODINGS: Final = ("utf-8-sig", "cp1251")
_KIND_BY_SUFFIX: Final = {
    ".docx": ImportKind.DOCX,
    ".xlsx": ImportKind.XLSX,
    ".xlsm": ImportKind.XLSX,
    ".csv": ImportKind.CSV,
    ".txt": ImportKind.CSV,
    ".zip": ImportKind.CSV,
}


def _suffix(name: str) -> str:
    return PurePosixPath(name.replace("\\", "/")).suffix.lower()


def _decode(content: bytes, name: str, warnings: list[str]) -> str:
    for encoding in _TEXT_ENCODINGS:
        try:
            text = content.decode(encoding)
        except UnicodeDecodeError:
            continue
        if encoding != _TEXT_ENCODINGS[0]:
            warnings.append(f"{name}: not UTF-8, decoded as {encoding}")
        return text
    raise ImportFormatError(f"{name}: cannot decode text (expected UTF-8)")  # pragma: no cover


def _trim(rows: Iterable[list[str]]) -> list[list[str]]:
    """Strip cells, drop trailing empty cells and fully empty rows."""
    out: list[list[str]] = []
    for row in rows:
        cells = [cell.strip() for cell in row]
        while cells and not cells[-1]:
            cells.pop()
        if cells:
            out.append(cells)
    return out


# --------------------------------------------------------------------------- docx


def read_docx(file: UploadedFile) -> RawUpload:
    try:
        document = docx.Document(io.BytesIO(file.content))
    except Exception as exc:  # python-docx raises several unrelated types for broken files
        raise ImportFormatError(f"{file.name}: not a valid .docx file ({exc})") from exc
    upload = RawUpload(ImportKind.DOCX, [file.name])
    for number, table in enumerate(document.tables, start=1):
        rows = _trim([cell.text for cell in row.cells] for row in table.rows)
        if rows:
            upload.tables.append(RawTable(f"{file.name}#table{number}", rows))
    upload.text = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    return upload


# --------------------------------------------------------------------------- xlsx


def xlsx_cell_text(value: object, number_format: str = "General") -> str:
    """Render an xlsx cell value as text (dates as ``dd.MM.yyyy``, percent cells x 100)."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, bool):
        return str(value).lower()
    if isinstance(value, int | float):
        number = float(value) * 100 if "%" in number_format else float(value)
        if number.is_integer():
            return str(int(number))
        return repr(round(number, 9))
    return str(value)


def read_xlsx(file: UploadedFile) -> RawUpload:
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(file.content), read_only=True, data_only=True)
    except Exception as exc:  # zipfile / openpyxl raise several unrelated types
        raise ImportFormatError(f"{file.name}: not a valid .xlsx file ({exc})") from exc
    upload = RawUpload(ImportKind.XLSX, [file.name])
    try:
        for sheet in workbook.worksheets:
            grid = (
                [xlsx_cell_text(c.value, getattr(c, "number_format", "General")) for c in row]
                for row in sheet.iter_rows()
            )
            rows = _trim(grid)
            if rows:
                upload.tables.append(RawTable(f"{file.name}#{sheet.title}", rows))
    finally:
        workbook.close()
    return upload


# --------------------------------------------------------------------------- csv / txt / zip


def read_csv_table(name: str, content: bytes, warnings: list[str]) -> RawTable:
    text = _decode(content, name, warnings)
    rows = _trim(csv.reader(io.StringIO(text, newline=""), delimiter=CSV_DELIMITER))
    return RawTable(name, rows)


def read_text(name: str, content: bytes, warnings: list[str]) -> list[str]:
    text = _decode(content, name, warnings)
    return [line.strip() for line in text.splitlines() if line.strip()]


def _unzip(file: UploadedFile) -> list[UploadedFile]:
    try:
        archive = zipfile.ZipFile(io.BytesIO(file.content))
    except zipfile.BadZipFile as exc:
        raise ImportFormatError(f"{file.name}: not a valid .zip file") from exc
    with archive:
        members = [
            info
            for info in archive.infolist()
            if not info.is_dir()
            and not PurePosixPath(info.filename).name.startswith(".")
            and "__MACOSX" not in info.filename
        ]
        if len(members) > MAX_ZIP_MEMBERS:
            raise ImportFormatError(f"{file.name}: too many files in the archive")
        if sum(info.file_size for info in members) > MAX_UNPACKED_BYTES:
            raise ImportFormatError(f"{file.name}: archive is too large when unpacked")
        return [
            UploadedFile(PurePosixPath(info.filename).name, archive.read(info))
            for info in sorted(members, key=lambda i: i.filename)
        ]


def read_csv_bundle(files: Sequence[UploadedFile]) -> RawUpload:
    upload = RawUpload(ImportKind.CSV, [f.name for f in files])
    expanded: list[UploadedFile] = []
    for file in files:
        expanded.extend(_unzip(file) if _suffix(file.name) == ".zip" else [file])
    for file in sorted(expanded, key=lambda f: f.name):
        suffix = _suffix(file.name)
        if suffix == ".csv":
            table = read_csv_table(file.name, file.content, upload.warnings)
            if table.rows:
                upload.tables.append(table)
        elif suffix == ".txt":
            upload.text.extend(read_text(file.name, file.content, upload.warnings))
        else:
            upload.warnings.append(f"{file.name}: ignored (only .csv and .txt are read)")
    return upload


# --------------------------------------------------------------------------- dispatch


def read_upload(files: Sequence[UploadedFile]) -> RawUpload:
    """Read an upload: one ``.docx``, one ``.xlsx``, or ``.csv``/``.txt``/``.zip`` files."""
    if not files:
        raise ImportFormatError("no files uploaded")
    kinds: set[ImportKind] = set()
    for file in files:
        kind = _KIND_BY_SUFFIX.get(_suffix(file.name))
        if kind is None:
            raise ImportFormatError(
                f"{file.name}: unsupported file type (expected .docx, .xlsx, .csv, .txt or .zip)"
            )
        kinds.add(kind)
    if len(kinds) > 1 or (kinds & {ImportKind.DOCX, ImportKind.XLSX} and len(files) > 1):
        raise ImportFormatError("upload one .docx, one .xlsx, or a set of .csv/.txt files")
    kind = kinds.pop()
    if kind is ImportKind.DOCX:
        return read_docx(files[0])
    if kind is ImportKind.XLSX:
        return read_xlsx(files[0])
    return read_csv_bundle(files)
