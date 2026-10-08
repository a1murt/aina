"""Import pipeline: readers, header recognition, values (FR-IMP-01), aliases and DQ-07
(FR-IMP-02), constraints from text (FR-IMP-03), template (FR-IMP-07)."""

from __future__ import annotations

import io
import zipfile
from datetime import date, datetime

import docx
import openpyxl
import pytest

from support import CASE_CSVS, CASE_DOCX, golden, golden_differences, require_case_docx
from twin_core.config import TwinConfig
from twin_core.importer import (
    ImportFormatError,
    ImportKind,
    RawTable,
    RawUpload,
    UploadedFile,
    build_import_report,
    build_template,
    configured_constraints,
    parse_constraints,
    parse_upload,
    read_upload,
    resolve_constraints,
    run_import,
)
from twin_core.importer.readers import xlsx_cell_text
from twin_core.importer.tables import classify_header, minutes_factor, recognize
from twin_core.importer.template import constraint_sentences
from twin_core.importer.values import normalize_header, parse_date, parse_int, parse_number

CASE_TEXT = [
    "Производство работает в 2 смены по 8 часов.",
    "Целевой показатель OEE - не менее 85%.",
    "Допустимый уровень брака - не более 2%.",
    "Максимально допустимый простой критического оборудования - 60 минут в сутки.",
    "План выпуска - не менее 5 500 автомобилей в месяц.",
]

LINES = [
    ["Дата", "Линия", "План", "Факт", "Время работы, ч", "Загрузка, %"],
    ["01.10.2026", "Сварка-1", "120", "118", "7.8", "98"],
    ["01.10.2026", "Окраска-1", "120", "115", "7.5", "94"],
    ["01.10.2026", "Сборка-1", "120", "121", "8.0", "100"],
]
DOWNTIME = [
    ["Дата", "Участок", "Оборудование", "Причина", "Длительность, мин"],
    ["01.10.2026", "Сварка", "ABB-01", "Ошибка датчика", "25"],
]
PLAN = [["Модель", "План на месяц"], ["Chevrolet Onix", "2500"]]
QUALITY = [
    ["Дата", "Участок", "Выпущено", "Брак", "% брака"],
    ["01.10.2026", "Сварка", "118", "2", "1,7"],
    ["01.10.2026", "Окраска", "115", "4", "3,5"],
    ["01.10.2026", "Сборка", "121", "1", "0,8"],
]


def _csv(name: str, rows: list[list[str]], encoding: str = "utf-8") -> UploadedFile:
    text = "\r\n".join(";".join(row) for row in rows) + "\r\n"
    return UploadedFile(name, text.encode(encoding))


def _bundle(**overrides: list[list[str]]) -> list[UploadedFile]:
    tables = {"lines": LINES, "downtime": DOWNTIME, "plan": PLAN, "quality": QUALITY}
    tables.update(overrides)
    return [_csv(f"{name}.csv", rows) for name, rows in tables.items()]


def _xlsx(sheets: dict[str, list[list[object]]]) -> UploadedFile:
    wb = openpyxl.Workbook()
    default = wb.active
    assert default is not None
    wb.remove(default)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return UploadedFile("book.xlsx", buffer.getvalue())


# --------------------------------------------------------------------------- values (FR-IMP-01)


@pytest.mark.parametrize(
    ("text", "value"),
    [
        ("1,7", 1.7),
        ("5 500", 5500.0),
        ("5 500", 5500.0),
        ("5 500,5", 5500.5),
        ("7.8", 7.8),
        ("98%", 98.0),
        ("1.234,5", 1234.5),
        ("1,234.5", 1234.5),
        ("-3", -3.0),
    ],
)
def test_parse_number(text: str, value: float) -> None:
    assert parse_number(text) == value


@pytest.mark.parametrize("text", ["", "abc", "1,2,3", "1.2.3"])
def test_parse_number_rejects(text: str) -> None:
    with pytest.raises(ValueError, match="not a number"):
        parse_number(text)


def test_parse_int_and_date() -> None:
    assert parse_int("120.0") == 120
    with pytest.raises(ValueError, match="whole"):
        parse_int("7,5")
    assert parse_date("01.10.2026") == date(2026, 10, 1)
    assert parse_date("1.10.2026") == date(2026, 10, 1)
    assert parse_date("2026-10-02") == date(2026, 10, 2)
    for bad in ("2026/10/01", "32.10.2026", "01.10.26"):
        with pytest.raises(ValueError, match=r"date|day is out of range"):
            parse_date(bad)
    assert normalize_header("  Время  РАБОТЫ, ч ") == "время работы, ч"


def test_xlsx_cell_text() -> None:
    assert xlsx_cell_text(None) == ""
    assert xlsx_cell_text(datetime(2026, 10, 1, 0, 0)) == "01.10.2026"  # noqa: DTZ001 (openpyxl is naive)
    assert xlsx_cell_text(date(2026, 10, 2)) == "02.10.2026"
    assert xlsx_cell_text(120) == "120"
    assert xlsx_cell_text(8.0) == "8"
    assert xlsx_cell_text(7.8) == "7.8"
    assert xlsx_cell_text(0.98, "0%") == "98"
    assert xlsx_cell_text(0.0348, "0.00%") == "3.48"
    assert xlsx_cell_text(True) == "true"
    assert xlsx_cell_text("Сварка-1") == "Сварка-1"


# --------------------------------------------------------------------------- readers


def test_unsupported_and_mixed_uploads() -> None:
    with pytest.raises(ImportFormatError, match="no files"):
        read_upload([])
    with pytest.raises(ImportFormatError, match="unsupported"):
        read_upload([UploadedFile("a.pdf", b"%PDF")])
    with pytest.raises(ImportFormatError, match="upload one"):
        read_upload([UploadedFile("a.docx", b""), UploadedFile("b.csv", b"")])
    with pytest.raises(ImportFormatError, match="upload one"):
        read_upload([UploadedFile("a.xlsx", b""), UploadedFile("b.xlsx", b"")])
    with pytest.raises(ImportFormatError, match=r"not a valid \.docx"):
        read_upload([UploadedFile("a.docx", b"garbage")])
    with pytest.raises(ImportFormatError, match=r"not a valid \.xlsx"):
        read_upload([UploadedFile("a.xlsx", b"garbage")])
    with pytest.raises(ImportFormatError, match=r"not a valid \.zip"):
        read_upload([UploadedFile("a.zip", b"garbage")])


def test_csv_bundle_reads_text_cp1251_and_ignores_others() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("readme.md", b"# ignored inside an archive")
    files = [
        *_bundle(),
        UploadedFile("notes.txt", "\n".join(CASE_TEXT).encode("cp1251")),
        UploadedFile("extra.zip", buffer.getvalue()),
    ]
    upload = read_upload(files)
    assert upload.kind is ImportKind.CSV
    assert len(upload.tables) == 4
    assert upload.text[0] == CASE_TEXT[0]
    assert any("cp1251" in w for w in upload.warnings)
    assert any("readme.md: ignored" in w for w in upload.warnings)


def test_zip_guards(monkeypatch: pytest.MonkeyPatch) -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("__MACOSX/._01.csv", b"junk")
        archive.writestr(".hidden.csv", b"junk")
        archive.writestr("dir/", b"")
        for path in CASE_CSVS:
            archive.write(path, path.name)
    upload = read_upload([UploadedFile("case.zip", buffer.getvalue())])
    assert [t.source for t in upload.tables] == [p.name for p in CASE_CSVS]
    monkeypatch.setattr("twin_core.importer.readers.MAX_ZIP_MEMBERS", 2)
    with pytest.raises(ImportFormatError, match="too many"):
        read_upload([UploadedFile("case.zip", buffer.getvalue())])
    monkeypatch.setattr("twin_core.importer.readers.MAX_ZIP_MEMBERS", 50)
    monkeypatch.setattr("twin_core.importer.readers.MAX_UNPACKED_BYTES", 10)
    with pytest.raises(ImportFormatError, match="too large"):
        read_upload([UploadedFile("case.zip", buffer.getvalue())])


# --------------------------------------------------------------------------- recognition


def test_classify_and_units() -> None:
    assert classify_header(["Дата", "ЛИНИЯ", "Факт"]) == "lines"
    assert classify_header(["Оборудование", "причина"]) == "downtime"
    assert classify_header(["Модель", "План"]) == "plan"
    assert classify_header(["Выпущено", "Брак"]) == "quality"
    assert classify_header(["Дата", "Линия"]) is None
    assert minutes_factor("Время работы, ч", default_hours=True) == 60
    assert minutes_factor("Время работы, мин", default_hours=True) == 1
    assert minutes_factor("Длительность (ч)", default_hours=False) == 60
    assert minutes_factor("Длительность", default_hours=False) == 1
    assert minutes_factor("Время работы", default_hours=True) == 60


def test_title_rows_reordered_columns_and_split_tables(cfg: TwinConfig) -> None:
    lines = [["Отчёт по линиям за октябрь"], [], *[[r[3], r[1], r[0], r[4], r[5]] for r in LINES]]
    upload = RawUpload(
        ImportKind.CSV,
        ["x"],
        tables=[
            RawTable("lines.csv", [row for row in lines if row]),
            RawTable("lines2.csv", [LINES[0], ["02.10.2026", "Сварка-1", "120", "111", "7.2"]]),
            RawTable("downtime.csv", DOWNTIME),
            RawTable("plan.csv", PLAN),
            RawTable("quality.csv", QUALITY),
            RawTable("notes", [["Целевой показатель OEE - не менее 80%."]]),
        ],
    )
    parsed = parse_upload(upload, cfg)
    assert [r.line for r in parsed.lines] == ["WELD-1", "PAINT-1", "ASSY-1", "WELD-1"]
    assert parsed.lines[0].plan_qty is None, "the reordered table has no plan column"
    assert parsed.lines[3].reported_load_pct is None
    assert parsed.text == ["Целевой показатель OEE - не менее 80%."]
    report = build_import_report(parsed, cfg, source="x")
    check = {c.key: c for c in report.constraint_checks}
    assert check["oee_target"].value == 0.8
    assert not check["oee_target"].matches
    assert check["oee_target"].source == "text"
    assert check["oee_target"].configured == 0.85
    assert any("quality: no row for 2026-10-02 WELD" in w for w in report.warnings)


def test_missing_tables_and_columns(cfg: TwinConfig) -> None:
    with pytest.raises(ImportFormatError, match="tables not found: quality") as info:
        parse_upload(read_upload(_bundle(quality=[["Пусто"]])), cfg)
    assert "missing table 'quality'" in info.value.problems
    no_fact_hours = [["Дата", "Линия", "Факт"], ["01.10.2026", "Сварка-1", "118"]]
    with pytest.raises(ImportFormatError, match="lacks columns: worked"):
        parse_upload(read_upload(_bundle(lines=no_fact_hours)), cfg)


def test_unknown_names_become_dq07_with_suggestion(cfg: TwinConfig) -> None:
    downtime = [
        *DOWNTIME,
        ["01.10.2026", "Окраска", "Камера-2", "Замена фильтра", "40"],
        ["01.10.2026", "Окраска", "Камера-02", "Замена фильтр", "40"],
        ["01.10.2026", "", "Конвейер-03", "Обрыв цепи", "сорок"],
        ["31.02.2026", "Сборка", "Конвейер-03", "Обрыв цепи", "5"],
        ["01.10.2026", "", "Конвейер-03", "Обрыв цепи", "0,5"],
    ]
    plan = [*PLAN, ["Lada Vesta", "100"]]
    parsed = parse_upload(read_upload(_bundle(downtime=downtime, plan=plan)), cfg)
    issues = [i.details for i in parsed.issues]
    assert {"kind": "equipment", "value": "Камера-2", "suggestion": "BOOTH-02"}.items() <= issues[
        0
    ].items()
    assert issues[0]["column"] == "Оборудование"
    assert issues[0]["row"] == 3
    assert issues[1]["kind"] == "reason"
    assert issues[1]["suggestion"] == "MT-FILTER"
    assert issues[2]["kind"] == "number"
    assert issues[2]["value"] == "сорок"
    assert issues[3]["kind"] == "date"
    assert issues[4]["kind"] == "product"
    assert parsed.issues[0].period_date == date(2026, 10, 1)
    assert parsed.issues[3].period_date is None
    assert [d.equipment for d in parsed.downtime] == ["ABB-01", "CONV-03"]
    assert parsed.downtime[1].area == "ASSY", "area derived from the equipment when empty"
    report = build_import_report(parsed, cfg, source="x")
    assert [i.rule_id for i in report.dq_issues].count("DQ-07") == 5


def test_inconsistent_rows_give_warnings(cfg: TwinConfig) -> None:
    lines = [
        *LINES,
        ["01.10.2026", "Сварка-1", "120", "117", "7.8", "98"],
        ["03.10.2026", "Сборка-1", "120", "5", "8", "100"],
    ]
    quality = [
        *QUALITY,
        ["01.10.2026", "Сварка", "118", "3", "2,5"],
        ["03.10.2026", "Сборка", "5", "9", "100"],
    ]
    downtime = [*DOWNTIME, ["01.10.2026", "Сборка", "ABB-01", "Ошибка датчика", "5"]]
    report = run_import(_bundle(lines=lines, quality=quality, downtime=downtime), cfg)
    text = "\n".join(report.warnings)
    assert "lines: duplicate row for 2026-10-01 WELD-1" in text
    assert "quality: duplicate row for 2026-10-01 WELD" in text
    assert "produced 118, line table says 117" in text
    assert "defects 9 exceed produced 5; capped" in text
    assert "2026-10-03 is not a working day" in text
    assert "area ASSY in the file, WELD in plant.yaml" in text


def test_no_usable_line_rows(cfg: TwinConfig) -> None:
    lines = [LINES[0], ["01.10.2026", "Линия-X", "120", "118", "7.8", "98"]]
    with pytest.raises(ImportFormatError, match="no usable rows") as info:
        run_import(_bundle(lines=lines), cfg)
    assert info.value.problems == ["Линия: 'Линия-X'"]


def test_xlsx_with_numbers_dates_and_notes_sheet(cfg: TwinConfig) -> None:
    gold = golden()
    lines: list[list[object]] = [[*LINES[0]]]
    for r in gold["shift_reports"]:
        day = date.fromisoformat(r["date"])
        name = cfg.lines[r["line"]].name_ru
        lines.append(
            [
                day,
                name,
                r["plan_qty"],
                r["produced_qty"],
                r["worked_min"] / 60,
                r["reported_load_pct"],
            ]
        )
    downtime: list[list[object]] = [[*DOWNTIME[0]]]
    names = {
        "ABB-01": "ABB-01",
        "ABB-04": "ABB-04",
        "BOOTH-02": "Камера-02",
        "CONV-03": "Конвейер-03",
    }
    for d in gold["downtime"]:
        downtime.append(
            [
                date.fromisoformat(d["date"]),
                cfg.areas[d["area"]].name_ru,
                names[d["equipment"]],
                d["reason_text_src"],
                d["duration_min"],
            ]
        )
    plan: list[list[object]] = [
        [*PLAN[0]],
        *[[r["model_src"], r["qty"]] for r in gold["plan"]["rows"]],
    ]
    quality: list[list[object]] = [[*QUALITY[0]]]
    for r in gold["shift_reports"]:
        quality.append(
            [
                date.fromisoformat(r["date"]),
                cfg.areas[r["area"]].name_ru,
                r["produced_qty"],
                r["defect_qty"],
                r["reported_defect_pct"],
            ]
        )
    book = _xlsx(
        {
            "Линии": lines,
            "Простои": downtime,
            "План": plan,
            "Качество": quality,
            "Вводные": [[t] for t in CASE_TEXT],
        }
    )
    report = run_import([book], cfg)
    assert golden_differences(report.to_json()) == []
    assert {c.source for c in report.constraint_checks} == {"text"}


# --------------------------------------------------------------------------- constraints


def test_parse_constraints_case_text() -> None:
    assert parse_constraints(CASE_TEXT) == golden()["constraints_parsed"]
    assert parse_constraints(["ничего"]) == {}
    assert parse_constraints(["OEE: не менее 82,5 %", "брака — не более 1,5%"]) == {
        "oee_target": 0.825,
        "defect_rate_limit": 0.015,
    }


def test_configured_constraints_and_resolution(cfg: TwinConfig) -> None:
    conf = configured_constraints(cfg, "2026-10")
    assert conf == golden()["constraints_parsed"]
    assert configured_constraints(cfg, "2027-01")["plant_target_per_month"] == (
        cfg.rules.thresholds.plant_target_per_month
    )
    effective, checks = resolve_constraints({"plant_target_per_month": 6000}, conf)
    assert effective["plant_target_per_month"] == 6000
    mismatch = [c for c in checks if not c.matches]
    assert [c.to_report() for c in mismatch] == [
        {
            "key": "plant_target_per_month",
            "value": 6000,
            "source": "text",
            "configured": 5500,
            "matches": False,
        }
    ]


def test_plan_target_from_text_drives_dq05(cfg: TwinConfig) -> None:
    files = [
        *_bundle(plan=[PLAN[0], ["Chevrolet Onix", "6000"]]),
        UploadedFile("n.txt", "не менее 6 000 автомобилей".encode()),
    ]
    report = run_import(files, cfg)
    assert report.plan.plant_target == 6000
    assert not [i for i in report.dq_issues if i.rule_id == "DQ-05"]


# --------------------------------------------------------------------------- template


def test_template_round_trip(cfg: TwinConfig) -> None:
    content = build_template(cfg)
    wb = openpyxl.load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["Линии", "Простои", "План", "Качество", "Вводные", "Справочник"]
    assert [c.value for c in wb["Линии"][1]] == LINES[0]
    upload = read_upload([UploadedFile("template.xlsx", content)])
    found, leftovers = recognize(upload.tables)
    assert set(found) == {"lines", "downtime", "plan", "quality"}
    assert all(not table.rows for table in found.values())
    assert parse_constraints(leftovers) == configured_constraints(cfg, None)
    assert constraint_sentences(cfg)[0] == CASE_TEXT[0]
    assert "5 500" in constraint_sentences(cfg)[-1]


def test_docx_reader_collects_tables_and_text(cfg: TwinConfig) -> None:
    require_case_docx()
    upload = read_upload([UploadedFile(CASE_DOCX.name, CASE_DOCX.read_bytes())])
    assert upload.kind is ImportKind.DOCX
    assert len(upload.tables) == 4
    assert upload.tables[0].source == "case2_data.docx#table1"
    assert "Дополнительные вводные" in upload.text
    document = docx.Document()
    document.add_paragraph("Только текст")
    buffer = io.BytesIO()
    document.save(buffer)
    with pytest.raises(ImportFormatError, match="tables not found"):
        run_import([UploadedFile("t.docx", buffer.getvalue())], cfg)
